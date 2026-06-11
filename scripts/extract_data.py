"""
extract_data.py

Regenerates the brochure data cache (data/product_data.json) from the source
Excel workbook. Run this when a new year's data lands, or when the product
file changes (new ranges, new colourways, corrected pricing, etc).

It reads the "Product Data" sheet, groups variant rows up to colourway level
(one entry per Product Code, with a sizes array), and derives the per range
metadata block the brochure uses for its navigation and range headers.

Usage:
  python3 extract_data.py /path/to/Think_Rugs_New_Products_XXXX.xlsx

Defaults to ./source/Think_Rugs_New_Products_2026.xlsx if no path is given.

This does NOT touch imagery. Image matching is by Product Code, so as long as
codes are stable, existing images in public/images keep working. After running
this, run npm run build.
"""

import sys, os, json
import openpyxl

ASSETS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data')
OUT = os.path.join(ASSETS, 'product_data.json')

def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None

def truthy(v):
    return str(v).strip().lower() in ('yes', 'y', 'true', '1') if v is not None else False

def main():
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(__file__), 'source', 'Think_Rugs_New_Products_2026.xlsx')
    if not os.path.isfile(src):
        print('Source workbook not found:', src)
        return

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb['Product Data'] if 'Product Data' in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    col = {name: i for i, name in enumerate(hdr)}

    def cell(row, name):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    products = {}   # keyed by Product Code (colourway level)
    order = []

    for row in rows:
        code = cell(row, 'Product Code')
        if not code:
            continue
        code = str(code).strip()

        # First word of Description = Range, second = Design (per build spec:
        # the source Range / Design columns are unreliable, parse Description).
        desc = str(cell(row, 'Description') or '').split()
        rng = desc[0] if len(desc) > 0 else (cell(row, 'Range') or '')
        dsg = desc[1] if len(desc) > 1 else (cell(row, 'Design') or '')

        size_entry = {
            'size': str(cell(row, 'Size (cm)') or '').strip(),
            'approx': cell(row, 'Approx Size'),
            'shape': cell(row, 'Shape'),
            'ws': num(cell(row, 'Wholesale Price (GBP)')),
            'rrp': num(cell(row, 'RRP (GBP)')),
            'vcode': str(cell(row, 'Variant Code') or '').strip(),
            'vdesc': str(cell(row, 'Description') or '').strip(),
        }

        if code not in products:
            order.append(code)
            feats = [cell(row, f) for f in ('Feature 1', 'Feature 2', 'Feature 3')]
            products[code] = {
                'code': code,
                'range': rng,
                'design': dsg,
                'c1': cell(row, 'Colour 1'),
                'c2': cell(row, 'Colour 2'),
                'style': cell(row, 'Style'),
                'construction': cell(row, 'Construction'),
                'm1': cell(row, 'Material 1'),
                'm2': cell(row, 'Material 2'),
                'room': cell(row, 'Room'),
                'wash': truthy(cell(row, 'Washable?')),
                'feats': [f for f in feats if f],
                'sizes': [],
            }
        products[code]['sizes'].append(size_entry)

    product_list = [products[c] for c in order]

    # ---- per range metadata ----
    ranges = {}
    for p in product_list:
        r = ranges.setdefault(p['range'], {
            'name': p['range'], 'colourways': 0, 'variants': 0,
            'designs': set(), 'styles': set(), 'construction': p['construction'],
            'm1': p['m1'], 'm2': p['m2'], 'wash': p['wash'], 'from': None, 'feats': [],
        })
        r['colourways'] += 1
        r['variants'] += len(p['sizes'])
        if p['design']:
            r['designs'].add(p['design'])
        if p['style']:
            r['styles'].add(p['style'])
        if p['wash']:
            r['wash'] = True
        if not r['feats'] and p['feats']:
            r['feats'] = p['feats']
        ws_vals = [s['ws'] for s in p['sizes'] if s['ws'] is not None]
        if ws_vals:
            lo = min(ws_vals)
            r['from'] = lo if r['from'] is None else min(r['from'], lo)

    range_list = []
    for name in sorted(ranges):
        r = ranges[name]
        r['designs'] = sorted(r['designs'])
        r['styles'] = sorted(r['styles'])
        range_list.append(r)

    json.dump({'products': product_list, 'ranges': range_list}, open(OUT, 'w'), indent=0)

    print(f'Wrote {OUT}')
    print(f'  {len(product_list)} colourways, '
          f'{sum(len(p["sizes"]) for p in product_list)} size variants, '
          f'{len(range_list)} ranges')
    print('Ranges:', ', '.join(r['name'] for r in range_list))
    print('Next step: npm run build')

if __name__ == '__main__':
    main()
